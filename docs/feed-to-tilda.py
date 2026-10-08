# -*- coding: utf-8 -*-
"""
Конвертер товарного фида 2YOUNG4U (выгрузка VipAvenue) → CSV для импорта в Tilda.

ЗАЧЕМ. Фид содержит 335 артикулов со всем, что нам нужно: название, описание,
состав, цвет, размеры, цена, остаток и до 10 прямых ссылок на фотографии.
Это снимает ручное заполнение полностью.

ЗАПУСК:
    python3 docs/feed-to-tilda.py <фид.xlsx> <выход.csv> [артикул ...]

    без списка артикулов — берутся ВСЕ;
    со списком — только указанные.

ПРАВИЛА ФОРМАТА (docs/csv-format.md):
  · колонки "Tilda UID" быть НЕ должно;
  · у РОДИТЕЛЯ пустые Price / Quantity / SKU / Editions;
  · у ВАРИАНТА пустые Category и характеристики;
  · Parent UID варианта = External ID родителя.

⚠️ ПРИ ИМПОРТЕ: «Группировка вариантов товара» → «Группировать по Parent UID».
    Сбрасывается между загрузками на «по SKU». Проверять каждый раз.
"""
import csv
import re
import sys
from collections import OrderedDict

import openpyxl

# ─────────────────── НАСТРОЙКИ ПО РЕШЕНИЯМ ЗАКАЗЧИЦЫ ───────────────────

QTY = 10              # кол-во на размер у всех (решение 07.10.2026)
SEASON_DEFAULT = "Всесезон"   # меняется на «Зима» / «Лето» точечно

# Сезон в фиде — это КОЛЛЕКЦИЯ (FW25/26, SS22), а не погода.
# Для фильтра нужна погода, поэтому коллекцию не переносим.
# Исключения по названию товара:
SEASON_BY_NAME = [
    ("winter tale", "Зима"),
]

# Категория фида  →  раздел в нашем каталоге.
# Левая часть сверяется по вхождению, первая подходящая выигрывает.
CAT_MAP = [
    ("Костюмы/Комплект",            "Комплекты"),
    ("Костюмы",                     "Костюмы"),
    ("Спортивная одежда/Олимпийки", "Олимпийки и толстовки"),
    ("Толстовки",                   "Олимпийки и толстовки"),
    ("Спортивная одежда/Спортивные брюки", "Брюки"),
    ("Брюки",                       "Брюки"),
    ("Футболки",                    "Футболки"),
    ("Шорты",                       "Шорты"),
    ("Пиджаки",                     "Мягкие пиджаки"),
    ("Верхняя одежда/Бомберы",      "Бомберы"),
    ("Платья",                      "Платья"),
    ("Юбки",                        "Юбки"),
]

GENDER_MAP = {"Мужская": "Мужское", "Женская": "Женское", "Унисекс": "Мужское"}

# Цвета приводим к единому написанию: в фиде «Черный», у нас «Чёрный».
COLOR_FIX = {"Черный": "Чёрный", "Кремовый": "Белый", "Зеленый": "Зелёный"}

COLOR_HEX = {
    "Синий": "#073763", "Белый": "#fffae5", "Чёрный": "#1d1d1d",
    "Серый": "#cccccc", "Бежевый": "#e8dcc8", "Голубой": "#9dc3e6",
    "Коричневый": "#6b4423", "Розовый": "#e8b4c8", "Зелёный": "#4a6741",
    "Жёлтый": "#f0d264", "Оранжевый": "#e08a3c", "Красный": "#c0392b",
    "Бордовый": "#7b241c", "Сиреневый": "#b8a9d4", "Серебряный": "#d5d8dc",
}

TRANSLIT = {
    'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z',
    'и':'i','й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r',
    'с':'s','т':'t','у':'u','ф':'f','х':'h','ц':'c','ч':'ch','ш':'sh','щ':'sch',
    'ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya',
}

HEADER = ["Brand", "SKU", "Mark", "Category", "Title", "Description", "Text",
          "Photo", "Price", "Quantity", "Price Old", "Editions", "Modifications",
          "External ID", "Parent UID",
          "Characteristics:Основной материал", "Characteristics:Подклад",
          "Characteristics:Цвет", "Characteristics:Сезон",
          "Characteristics:Стиль", "Characteristics:Капсула",
          "Weight", "Length", "Width", "Height", "Url"]

SIZE_ORDER = ["XS", "S", "M", "L", "XL", "XXL", "2XL", "3XL", "XXXL"]


def slugify(text, prefix=""):
    t = (text or "").lower()
    out = "".join(TRANSLIT.get(ch, ch if ch.isalnum() else "-") for ch in t)
    out = re.sub(r"-+", "-", out).strip("-")
    return (prefix + out)[:60].strip("-")


def material_of(sostav):
    """Короткое значение для фильтра — первый материал из состава.
    Весь состав целиком в фильтр класть нельзя: у каждого товара получится
    своё уникальное значение и фильтр станет бесполезным
    (см. docs/product-fields.md)."""
    if not sostav:
        return ""
    m = re.search(r"[а-яё]+", str(sostav).lower())
    if not m:
        return ""
    word = m.group(0)
    return word[0].upper() + word[1:]


def clean_descr(text):
    """В фиде описание идёт списком через дефисы в начале строк."""
    if not text:
        return ""
    parts = [p.strip(" -–—\n\r\t") for p in re.split(r"\s*-(?=[А-ЯA-Z])", str(text))]
    parts = [p for p in parts if p]
    return " ".join(parts)


# У женского раздела подраздел называется короче — «Олимпийки».
SUB_BY_GENDER = {("Женское", "Олимпийки и толстовки"): "Олимпийки"}


def map_category(feed_cat, gender):
    razdel = GENDER_MAP.get(gender, "Мужское")
    sub = None
    for needle, name in CAT_MAP:
        if needle.lower() in (feed_cat or "").lower():
            sub = name
            break
    if not sub:
        return razdel
    sub = SUB_BY_GENDER.get((razdel, sub), sub)
    return "%s;%s>>>%s" % (razdel, razdel, sub)


def size_key(s):
    s = (s or "").upper()
    return (SIZE_ORDER.index(s) if s in SIZE_ORDER else 99, s)


def read_feed(path, only=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Список товаров"]
    H = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    get = lambda r, name: ws.cell(row=r, column=H[name]).value

    products = OrderedDict()
    for r in range(2, ws.max_row + 1):
        art = get(r, "Артикул")
        if art in (None, ""):
            continue
        art = str(art).strip()
        if only and art not in only:
            continue

        p = products.setdefault(art, {"sizes": [], "photos": []})
        p["sku"] = art
        p["title"] = (get(r, "Название товара") or "").strip()
        p["gender"] = get(r, "Пол")
        p["feed_cat"] = get(r, "Категория товара")
        p["price"] = get(r, "Цена")
        p["sostav"] = get(r, "Состав")
        p["descr"] = get(r, "Описание товара")
        color = (get(r, "Цвет") or "").strip()
        p["color"] = COLOR_FIX.get(color, color)

        if not p["photos"]:
            for i in range(1, 11):
                u = get(r, "Фото%d" % i)
                if u:
                    p["photos"].append(str(u).strip())

        size = (get(r, "Размер производителя") or "").strip()
        if size and size not in p["sizes"]:
            p["sizes"].append(size)
    return products


def build_rows(products):
    rows = []
    for art, p in products.items():
        title = p["title"]
        color = p["color"]
        slug = slugify("%s %s" % (title, color), prefix="")
        season = SEASON_DEFAULT
        for needle, value in SEASON_BY_NAME:
            if needle in title.lower():
                season = value

        price = "%.2f" % float(p["price"] or 0)

        parent = {k: "" for k in HEADER}
        parent["Brand"] = "2young4u"
        parent["Category"] = map_category(p["feed_cat"], p["gender"])
        parent["Title"] = title
        parent["Text"] = clean_descr(p["descr"])
        parent["Photo"] = " ".join(p["photos"])
        parent["External ID"] = slug
        parent["Characteristics:Основной материал"] = material_of(p["sostav"])
        parent["Characteristics:Подклад"] = "Без подклада"
        parent["Characteristics:Цвет"] = color
        parent["Characteristics:Сезон"] = season
        for d in ("Weight", "Length", "Width", "Height"):
            parent[d] = "0"
        parent["Url"] = "https://2young4u.ru/p/" + slug
        rows.append(parent)

        hexv = COLOR_HEX.get(color, "#cccccc")
        for size in sorted(p["sizes"], key=size_key):
            v = {k: "" for k in HEADER}
            v["SKU"] = art
            v["Title"] = "%s - %s - %s" % (title, size, color)
            v["Price"] = price
            v["Quantity"] = str(QTY)
            v["Editions"] = "Размер:%s;Цвет:%s %s" % (size, color, hexv)
            v["External ID"] = "%s-%s" % (slug, size.lower())
            v["Parent UID"] = slug
            for d in ("Weight", "Length", "Width", "Height"):
                v[d] = "0"
            rows.append(v)
    return rows


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return
    feed, out = sys.argv[1], sys.argv[2]
    only = set(sys.argv[3:]) or None
    products = read_feed(feed, only)
    rows = build_rows(products)
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HEADER, delimiter=";",
                           quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print("%s — товаров: %d, строк: %d" % (out, len(products), len(rows)))


if __name__ == "__main__":
    main()
