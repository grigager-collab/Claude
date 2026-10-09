# -*- coding: utf-8 -*-
"""
Генератор блоков для СТРАНИЦ КАРТОЧЕК из товарного фида 2YOUNG4U.

На каждой странице товара нужен блок T123 с текстами аккордеонов и атрибутами
хлебных крошек. Раньше он набирался руками. Этот скрипт собирает его из фида:
описание и состав там уже есть.

ЗАПУСК:
    python3 docs/gen-page-blocks.py <фид.xlsx> <выход.md> [артикул ...]

ЧТО ПОЛУЧАЕТСЯ — на каждый товар готовый кусок для вставки:

    <div class="t2y-acc-source"
         data-crumb-parent="Мужское"  data-crumb-parent-url="/men"
         data-crumb-name="Олимпийки и толстовки" data-crumb-url="/men/hoodies">
      <div data-title="Детали">…</div>
      <div data-title="Ткань, состав, уход">…</div>
      <div data-title="Параметры модели на фото">…</div>
    </div>

Движок аккордеонов и скрипт крошек живут в HEAD, здесь только содержимое.
"""
import re
import sys

import openpyxl

# Раздел → адрес, для атрибутов крошек
URLS = {
    "Мужское": "/men", "Женское": "/women",
    "Костюмы": "/suits", "Олимпийки и толстовки": "/hoodies",
    "Олимпийки": "/hoodies", "Мягкие пиджаки": "/blazers",
    "Брюки": "/trousers", "Футболки": "/tshirts",
    "Комплекты": "/sets", "Шорты": "/shorts",
}

CAT_MAP = [
    ("Костюмы/Комплект", "Комплекты"),
    ("Костюмы", "Костюмы"),
    ("Спортивная одежда/Олимпийки", "Олимпийки и толстовки"),
    ("Толстовки", "Олимпийки и толстовки"),
    ("Спортивная одежда/Спортивные брюки", "Брюки"),
    ("Брюки", "Брюки"),
    ("Футболки", "Футболки"),
    ("Шорты", "Шорты"),
    ("Пиджаки", "Мягкие пиджаки"),
]

GENDER = {"Мужская": "Мужское", "Женская": "Женское", "Унисекс": "Мужское"}

# Транслитерация и приведение цвета — ОДИН В ОДИН как в feed-to-tilda.py,
# иначе адреса в блоках и в CSV разойдутся.
COLOR_FIX = {"Черный": "Чёрный", "Кремовый": "Белый", "Зеленый": "Зелёный"}

TRANSLIT = {
    'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z',
    'и':'i','й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r',
    'с':'s','т':'t','у':'u','ф':'f','х':'h','ц':'c','ч':'ch','ш':'sh','щ':'sch',
    'ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya',
}


def slugify(text):
    t = (text or "").lower()
    out = "".join(TRANSLIT.get(ch, ch if ch.isalnum() else "-") for ch in t)
    out = re.sub(r"-+", "-", out).strip("-")
    return out[:60].strip("-")

# Уход зависит от состава: шерсть и лён просят бережнее, чем хлопок.
CARE_COTTON = [
    "Стирать при температуре до 30 °C на бережном режиме.",
    "Использовать мягкие средства для цветных тканей.",
    "Не отбеливать, не использовать агрессивные пятновыводители.",
    "Сушить в расправленном виде вдали от прямых солнечных лучей и батарей.",
    "Гладить с изнаночной стороны при средней температуре.",
]
CARE_WOOL = [
    "Рекомендована сухая чистка.",
    "Ручная стирка при температуре до 30 °C со средством для шерсти.",
    "Не замачивать, не выкручивать, не отбеливать.",
    "Сушить в расправленном виде на горизонтальной поверхности.",
    "Гладить через влажную ткань при низкой температуре.",
]
CARE_LINEN = [
    "Стирать при температуре до 30 °C на деликатном режиме.",
    "Не отбеливать. Отжим на малых оборотах.",
    "Сушить в расправленном виде, лён легко заминается.",
    "Гладить слегка влажным при высокой температуре.",
]


def care_for(sostav):
    s = (sostav or "").lower()
    if "шерст" in s or "кашемир" in s:
        return CARE_WOOL
    if "лен" in s or "лён" in s:
        return CARE_LINEN
    return CARE_COTTON


def esc(t):
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def split_descr(text):
    """В фиде описание — список через дефисы в начале предложений."""
    if not text:
        return []
    parts = re.split(r"\s*-(?=[А-ЯA-Z])", str(text))
    return [p.strip(" -–—\n\r\t") for p in parts if p.strip(" -–—\n\r\t")]


def map_cat(feed_cat, gender):
    razdel = GENDER.get(gender, "Мужское")
    for needle, name in CAT_MAP:
        if needle.lower() in (feed_cat or "").lower():
            return razdel, name
    return razdel, None


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return
    feed, out = sys.argv[1], sys.argv[2]
    only = set(sys.argv[3:]) or None

    wb = openpyxl.load_workbook(feed, data_only=True)
    ws = wb["Список товаров"]
    H = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    g = lambda r, n: ws.cell(row=r, column=H[n]).value

    prod = {}
    for r in range(2, ws.max_row + 1):
        a = g(r, "Артикул")
        if a in (None, ""):
            continue
        a = str(a).strip()
        if only and a not in only:
            continue
        # ⚠️ Название берём ПОСЛЕДНЕЕ по файлу, как и feed-to-tilda.py.
        # В фиде один артикул может быть записан под ДВУМЯ названиями
        # (например 21260209: «Олимпийка Space реглан» и «Олимпийка хлопковая»).
        # Если скрипты возьмут разные — адреса в блоках и в CSV разойдутся.
        color = (g(r, "Цвет") or "").strip()
        prod[a] = {
            "name": (g(r, "Название товара") or "").strip(),
            "descr": g(r, "Описание товара"),
            "sostav": g(r, "Состав"),
            "cat": g(r, "Категория товара"),
            "gender": g(r, "Пол"),
            "country": g(r, "Страна производителя"),
            "color": COLOR_FIX.get(color, color),
        }

    # Файл отдаём ЧИСТЫМ HTML, без markdown-разметки: если скопировать лишнее,
    # оно окажется внутри комментария и на странице не отобразится.
    # (08.10.2026: markdown-версия привела к тому, что на карточку попали
    #  символы ```html и ---, и Tilda вывела их текстом.)
    lines = [
        "<!-- ================================================================",
        "     БЛОКИ ДЛЯ СТРАНИЦ КАРТОЧЕК — собрано из товарного фида 2YOUNG4U.",
        "",
        "     КАК ПОЛЬЗОВАТЬСЯ: на странице карточки добавить блок",
        "     «Другое → HTML-код» и вставить нужный кусок.",
        "",
        "     ИСКАТЬ ПО АДРЕСУ СТРАНИЦЫ — он первой строкой в комментарии,",
        "     вида /p/olimpiyka-hlopkovaya-siniy. По названию искать НЕ надо:",
        "     в фиде один артикул бывает записан под двумя названиями.",
        "     Движок аккордеонов и крошек уже лежит в HEAD — здесь только тексты.",
        "",
        "     КОПИРОВАТЬ можно с запасом: все пояснения внутри комментариев",
        "     и на странице не видны.",
        "",
        "     ⚠️ «Параметры модели на фото» фид не содержит — там прочерки,",
        "     заменить реальными данными со съёмки.",
        "     ================================================================ -->",
        "",
    ]

    for art, p in prod.items():
        razdel, cat = map_cat(p["cat"], p["gender"])
        crumb_url = ""
        if cat:
            crumb_url = URLS.get(razdel, "") + URLS.get(cat, "")

        details = split_descr(p["descr"])
        care = care_for(p["sostav"])

        lines.append("")
        slug = slugify("%s %s" % (p["name"], p.get("color", "")))
        lines.append("<!-- ====== /p/%s ======" % slug)
        lines.append("         %s · артикул %s ====== -->" % (p["name"], art))
        lines.append('<div class="t2y-acc-source"')
        lines.append('     data-crumb-parent="%s" data-crumb-parent-url="%s"'
                     % (razdel, URLS.get(razdel, "/")))
        if cat:
            lines.append('     data-crumb-name="%s" data-crumb-url="%s">'
                         % (cat, crumb_url))
        else:
            lines.append('     data-crumb-name="" data-crumb-url="">')

        lines.append('  <div data-title="Детали">')
        for d in details:
            lines.append("    <p>%s</p>" % esc(d))
        if p.get("country"):
            lines.append("    <p>Страна производства: %s.</p>"
                         % esc(str(p["country"]).capitalize()))
        lines.append("  </div>")

        lines.append('  <div data-title="Ткань, состав, уход">')
        if p["sostav"]:
            lines.append("    <p>Состав: %s.</p>" % esc(p["sostav"]))
        for c in care:
            lines.append("    <p>%s</p>" % esc(c))
        lines.append("  </div>")

        lines.append('  <div data-title="Параметры модели на фото">')
        lines.append("    <p>Рост модели: — см.</p>")
        lines.append("    <p>На модели размер: —.</p>")
        lines.append("  </div>")
        lines.append("</div>")

    open(out, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    print("%s — блоков: %d" % (out, len(prod)))


if __name__ == "__main__":
    main()
